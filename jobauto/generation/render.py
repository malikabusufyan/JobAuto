from datetime import date
from pathlib import Path

from jinja2 import Environment, FileSystemLoader

TEMPLATES_DIR = Path(__file__).parent / "templates"
_env = Environment(loader=FileSystemLoader(str(TEMPLATES_DIR)), autoescape=True)


def _html_to_pdf(html: str, out_path: Path):
    from playwright.sync_api import sync_playwright

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()
        page.set_content(html, wait_until="load")
        page.pdf(path=str(out_path), print_background=True, format="Letter")
        browser.close()


def render_resume_pdf(profile: dict, tailored: dict, out_path: Path):
    template = _env.get_template("resume.html")
    html = template.render(
        personal=profile["personal"],
        headline=tailored.get("headline", ""),
        summary=tailored.get("summary", ""),
        skills=tailored.get("skills", []),
        experience=tailored.get("experience", []),
        projects=tailored.get("projects", []),
        education=tailored.get("education", profile.get("education", [])),
        certifications=profile.get("certifications", []),
        achievements=profile.get("achievements", []),
    )
    _html_to_pdf(html, out_path)


def render_resume_docx(profile: dict, tailored: dict, out_path: Path):
    import docx
    from docx.shared import Pt

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()
    personal = profile["personal"]

    title = doc.add_heading(personal["full_name"], level=0)
    doc.add_paragraph(tailored.get("headline", ""))
    contact = " | ".join(
        filter(None, [personal.get("email"), personal.get("phone"), personal.get("location"),
                      personal.get("linkedin"), personal.get("github")])
    )
    doc.add_paragraph(contact)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(tailored.get("summary", ""))

    doc.add_heading("Skills", level=1)
    doc.add_paragraph(" | ".join(tailored.get("skills", [])))

    doc.add_heading("Experience", level=1)
    for job in tailored.get("experience", []):
        p = doc.add_paragraph()
        run = p.add_run(f"{job.get('title', '')}, {job.get('company', '')}")
        run.bold = True
        p.add_run(f"    {job.get('start_date', '')} - {job.get('end_date', '')}").italic = True
        for bullet in job.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")

    if tailored.get("projects"):
        doc.add_heading("Projects", level=1)
        for proj in tailored["projects"]:
            p = doc.add_paragraph()
            p.add_run(proj.get("name", "")).bold = True
            if proj.get("description"):
                doc.add_paragraph(proj["description"])
            for bullet in proj.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("Education", level=1)
    for edu in tailored.get("education", profile.get("education", [])):
        p = doc.add_paragraph()
        p.add_run(f"{edu.get('degree', '')} {edu.get('field', '')} - {edu.get('institution', '')}").bold = True
        p.add_run(f"    {edu.get('start_date', '')} - {edu.get('end_date', '')}").italic = True

    if profile.get("achievements"):
        doc.add_heading("Achievements", level=1)
        for achievement in profile["achievements"]:
            doc.add_paragraph(achievement, style="List Bullet")

    for style_name in ["Normal"]:
        style = doc.styles[style_name]
        style.font.size = Pt(10.5)

    doc.save(str(out_path))


def render_cover_letter_pdf(profile: dict, job: dict, letter_text: str, out_path: Path):
    template = _env.get_template("cover_letter.html")
    paragraphs = [p.strip() for p in letter_text.split("\n\n") if p.strip()]
    html = template.render(
        personal=profile["personal"],
        date=date.today().strftime("%B %d, %Y"),
        job_title=job.get("title", ""),
        company=job.get("company", ""),
        paragraphs=paragraphs,
    )
    _html_to_pdf(html, out_path)
